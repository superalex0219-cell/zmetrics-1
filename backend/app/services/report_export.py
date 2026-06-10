"""Report export: collect the full report bundle and render it to a format.

REPORT-X: PDF (WeasyPrint + HTML template), DOCX (python-docx), XLSX (openpyxl),
CSV and JSON. Renderers are pure functions ``(data, images) -> bytes`` so they
are unit-testable without a database.

SAFETY (safety.md): every format displays the analysis method prominently —
mock results carry the «⚠ Синтетические данные» badge, real-CV results show
the model version; recommendations are always rendered with their review
status (never as approved facts) and parameter suggestions are reference-only.
"""

from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

EXPORT_FORMATS = ("json", "csv", "xlsx", "docx", "pdf")

MEDIA_TYPES = {
    "json": "application/json",
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}

MOCK_BADGE = "⚠ Синтетические данные (mock-пайплайн) — не для производственных решений"

RECOMMENDATION_STATUS_RU = {
    "requires_human_review": "ТРЕБУЕТ ПРОВЕРКИ",
    "reviewed": "рассмотрена",
    "accepted": "принята",
    "rejected": "отклонена",
}

PASSPORT_FIELDS_RU: list[tuple[str, str]] = [
    ("revision_number", "Ревизия"),
    ("status", "Статус"),
    ("explosive_type", "Тип ВВ"),
    ("total_explosive_kg", "Общий заряд, кг"),
    ("number_of_holes", "Кол-во скважин"),
    ("hole_diameter_mm", "Диаметр скважины, мм"),
    ("hole_depth_m", "Глубина скважины, м"),
    ("burden_m", "ЛНС (burden), м"),
    ("spacing_m", "Расстояние в ряду (spacing), м"),
    ("stemming_m", "Забойка, м"),
    ("target_p80_mm", "Целевой P80, мм"),
    ("notes", "Примечания"),
]

RESULT_FIELDS_RU: list[tuple[str, str]] = [
    ("p10_mm", "P10, мм"),
    ("p50_mm", "P50, мм"),
    ("p80_mm", "P80, мм"),
    ("rosin_rammler_n", "Rosin-Rammler n"),
    ("rosin_rammler_xc", "Rosin-Rammler xc, мм"),
    ("uniformity_index", "Индекс однородности"),
    ("oversize_percent", "Негабарит, %"),
    ("fines_percent", "Переизмельчение, %"),
    ("total_particles_counted", "Камней измерено"),
    ("total_volume_m3", "Суммарный объём, м³"),
    ("confidence_score", "Confidence"),
]


def _num(value: Any) -> float | None:
    return float(value) if value is not None else None


def _fmt(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


async def collect_report_data(db: AsyncSession, report_id: UUID) -> dict | None:
    """Full export bundle for one report; None when the report does not exist."""
    from app.db.models.analysis import AnalysisJob, AnalysisResult, ModelVersion
    from app.db.models.blast import BlastEvent, Device
    from app.db.models.capture import CaptureSession
    from app.db.models.passport import BlastPassport
    from app.db.models.quarry import Quarry, SiteSection
    from app.db.models.report import Recommendation, Report

    report = (await db.execute(select(Report).where(Report.id == report_id))).scalar_one_or_none()
    if report is None:
        return None

    ar = (await db.execute(
        select(AnalysisResult).where(AnalysisResult.id == report.analysis_result_id)
    )).scalar_one_or_none()

    job = session = blast = passport = section = quarry = device = model = None
    if ar is not None:
        job = (await db.execute(
            select(AnalysisJob).where(AnalysisJob.id == ar.job_id)
        )).scalar_one_or_none()
    if job is not None:
        if job.model_version_id is not None:
            model = (await db.execute(
                select(ModelVersion).where(ModelVersion.id == job.model_version_id)
            )).scalar_one_or_none()
        session = (await db.execute(
            select(CaptureSession).where(CaptureSession.id == job.capture_session_id)
        )).scalar_one_or_none()
    if session is not None:
        device = (await db.execute(
            select(Device).where(Device.id == session.device_id)
        )).scalar_one_or_none()
        blast = (await db.execute(
            select(BlastEvent).where(BlastEvent.id == session.blast_event_id)
        )).scalar_one_or_none()
    if blast is not None:
        passport = (await db.execute(
            select(BlastPassport).where(BlastPassport.id == blast.passport_id)
        )).scalar_one_or_none()
    if passport is not None:
        section = (await db.execute(
            select(SiteSection).where(SiteSection.id == passport.site_section_id)
        )).scalar_one_or_none()
    if section is not None:
        quarry = (await db.execute(
            select(Quarry).where(Quarry.id == section.quarry_id)
        )).scalar_one_or_none()

    recs = list((await db.execute(
        select(Recommendation).where(Recommendation.report_id == report_id)
        .order_by(Recommendation.created_at)
    )).scalars().all())

    model_type = model.model_type if model is not None else None
    analysis_method = "real" if (model_type and model_type != "mock") else "mock"

    return {
        "exported_at": datetime.now(tz=timezone.utc).isoformat(timespec="seconds"),
        "method": {
            "analysis_method": analysis_method,
            "model_version_tag": model.version_tag if model is not None else None,
            "label": (
                MOCK_BADGE if analysis_method == "mock"
                else f"Реальный CV-пайплайн (модель {model.version_tag})"
            ),
        },
        "report": {
            "id": str(report.id),
            "title": report.title,
            "report_type": report.report_type,
            "created_at": report.created_at.isoformat(),
        },
        "quarry": {
            "name": quarry.name if quarry else None,
            "location": quarry.location_description if quarry else None,
            "section": section.name if section else None,
            "block_number": section.block_number if section else None,
        },
        "blast_event": {
            "blast_datetime": blast.blast_datetime.isoformat() if blast else None,
            "actual_explosive_kg": _num(blast.actual_explosive_kg) if blast else None,
            "weather_conditions": blast.weather_conditions if blast else None,
            "notes": blast.notes if blast else None,
        },
        "passport": {
            "revision_number": passport.revision_number if passport else None,
            "status": passport.status.value if passport else None,
            "explosive_type": passport.explosive_type if passport else None,
            "total_explosive_kg": _num(passport.total_explosive_kg) if passport else None,
            "number_of_holes": passport.number_of_holes if passport else None,
            "hole_diameter_mm": _num(passport.hole_diameter_mm) if passport else None,
            "hole_depth_m": _num(passport.hole_depth_m) if passport else None,
            "burden_m": _num(passport.burden_m) if passport else None,
            "spacing_m": _num(passport.spacing_m) if passport else None,
            "stemming_m": _num(passport.stemming_m) if passport else None,
            "target_p80_mm": _num(passport.target_p80_mm) if passport else None,
            "notes": passport.notes if passport else None,
        },
        "capture": {
            "session_id": str(session.id) if session else None,
            "capture_datetime": session.capture_datetime.isoformat() if session else None,
            "device": f"{device.model} ({device.serial_number})" if device else None,
            "frame_index": getattr(job, "frame_index", None) if job else None,
        },
        "analysis_result": {
            "p10_mm": _num(ar.p10_mm) if ar else None,
            "p50_mm": _num(ar.p50_mm) if ar else None,
            "p80_mm": _num(ar.p80_mm) if ar else None,
            "rosin_rammler_n": _num(ar.rosin_rammler_n) if ar else None,
            "rosin_rammler_xc": _num(ar.rosin_rammler_xc) if ar else None,
            "uniformity_index": _num(ar.uniformity_index) if ar else None,
            "oversize_percent": _num(ar.oversize_percent) if ar else None,
            "fines_percent": _num(ar.fines_percent) if ar else None,
            "total_particles_counted": ar.total_particles_counted if ar else None,
            "total_volume_m3": _num(ar.total_volume_m3) if ar else None,
            "confidence_score": _num(ar.confidence_score) if ar else None,
            "confidence_notes": ar.confidence_notes if ar else None,
            "size_distribution": (ar.size_distribution or []) if ar else [],
        },
        "recommendations": [
            {
                "id": str(r.id),
                "status": r.status.value,
                "status_ru": RECOMMENDATION_STATUS_RU.get(r.status.value, r.status.value),
                "recommendation_text": r.recommendation_text,
                "parameter_suggestions": r.parameter_suggestions,
                "reviewed_at": r.reviewed_at.isoformat() if r.reviewed_at else None,
                "reviewer_notes": r.reviewer_notes,
            }
            for r in recs
        ],
        # MinIO prefix of the job's pipeline artifacts (visuals live here)
        "_pipeline_prefix": _job_prefix(job) if job is not None else None,
        "_capture_session_id": str(job.capture_session_id) if job is not None else None,
        "_frame_index": int(getattr(job, "frame_index", 0) or 0) if job is not None else 0,
    }


def _job_prefix(job: Any) -> str:
    base = f"sessions/{job.capture_session_id}/pipeline"
    idx = int(getattr(job, "frame_index", 0) or 0)
    return base if idx == 0 else f"{base}/f{idx:04d}"


async def collect_report_images(db: AsyncSession, storage: Any, data: dict) -> dict[str, bytes]:
    """Best-effort fetch of the report visuals: left frame, mask overlay, depth map."""
    from app.db.models.artifact import Artifact, ArtifactType

    images: dict[str, bytes] = {}
    prefix = data.get("_pipeline_prefix")
    session_id = data.get("_capture_session_id")
    if not prefix or not session_id:
        return images

    def _try_fetch(name: str, bucket: str, key: str) -> None:
        try:
            images[name] = storage.download_bytes(bucket, key)
        except Exception:
            pass  # отсутствие картинки — не ошибка экспорта

    idx = int(data.get("_frame_index") or 0)
    frame_filter = Artifact.frame_index == idx
    if idx == 0:
        frame_filter = frame_filter | Artifact.frame_index.is_(None)
    frame = (await db.execute(
        select(Artifact).where(
            Artifact.capture_session_id == UUID(session_id),
            Artifact.artifact_type == ArtifactType.LEFT_FRAME,
            frame_filter,
        ).order_by(Artifact.frame_index).limit(1)
    )).scalar_one_or_none()
    if frame is not None:
        _try_fetch("left_frame", frame.storage_bucket, frame.storage_key)

    _try_fetch("mask_overlay", "zmetrics-artifacts", f"{prefix}/report_visuals/mask_overlay.png")
    _try_fetch("depth_color", "zmetrics-artifacts", f"{prefix}/report_visuals/depth_color.png")
    return images


def _public(data: dict) -> dict:
    return {k: v for k, v in data.items() if not k.startswith("_")}


# ── Renderers ────────────────────────────────────────────────────────────────


def render_json(data: dict) -> bytes:
    return json.dumps(_public(data), ensure_ascii=False, indent=2).encode("utf-8")


def render_csv(data: dict) -> bytes:
    """Size distribution as CSV (utf-8-sig so Excel opens Cyrillic correctly)."""
    import csv

    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(["Отчёт", data["report"]["title"]])
    writer.writerow(["Метод анализа", data["method"]["label"]])
    writer.writerow(["Карьер", _fmt(data["quarry"]["name"])])
    writer.writerow(["Взрыв", _fmt(data["blast_event"]["blast_datetime"])])
    writer.writerow([])
    for key, label in RESULT_FIELDS_RU:
        writer.writerow([label, _fmt(data["analysis_result"].get(key))])
    writer.writerow([])
    writer.writerow(["Размер, мм", "Прохождение, %"])
    for row in data["analysis_result"]["size_distribution"]:
        writer.writerow([row.get("size_mm"), row.get("cumulative_passing_pct")])
    return buf.getvalue().encode("utf-8-sig")


def render_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)
    red_bold = Font(bold=True, color="FFAA0000")

    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Отчёт", data["report"]["title"]])
    ws.append(["Метод анализа", data["method"]["label"]])
    ws["B2"].font = red_bold if data["method"]["analysis_method"] == "mock" else bold
    ws.append(["Сформирован", data["exported_at"]])
    ws.append(["Карьер", _fmt(data["quarry"]["name"])])
    ws.append(["Участок", _fmt(data["quarry"]["section"])])
    ws.append(["Взрыв", _fmt(data["blast_event"]["blast_datetime"])])
    ws.append(["Устройство", _fmt(data["capture"]["device"])])
    ws.append([])
    for key, label in RESULT_FIELDS_RU:
        ws.append([label, _fmt(data["analysis_result"].get(key))])
    ws.append(["Примечания", _fmt(data["analysis_result"]["confidence_notes"])])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70

    ws2 = wb.create_sheet("Грансостав")
    ws2.append(["Размер, мм", "Прохождение, %"])
    ws2["A1"].font = ws2["B1"].font = bold
    for row in data["analysis_result"]["size_distribution"]:
        ws2.append([row.get("size_mm"), row.get("cumulative_passing_pct")])

    ws3 = wb.create_sheet("Паспорт БВР")
    for key, label in PASSPORT_FIELDS_RU:
        ws3.append([label, _fmt(data["passport"].get(key))])
    ws3.column_dimensions["A"].width = 34

    ws4 = wb.create_sheet("Рекомендации")
    ws4.append(["Статус", "Текст", "Параметры (справочно)", "Заметки проверяющего"])
    for cell in ws4[1]:
        cell.font = bold
    for rec in data["recommendations"]:
        ws4.append([
            rec["status_ru"],
            rec["recommendation_text"],
            json.dumps(rec["parameter_suggestions"], ensure_ascii=False)
            if rec["parameter_suggestions"] else "—",
            _fmt(rec["reviewer_notes"]),
        ])
    ws4.column_dimensions["B"].width = 90

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_docx(data: dict, images: dict[str, bytes] | None = None) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    images = images or {}
    doc = Document()
    doc.add_heading(data["report"]["title"], level=1)

    method = doc.add_paragraph()
    run = method.add_run(data["method"]["label"])
    run.bold = True
    run.font.size = Pt(12)
    if data["method"]["analysis_method"] == "mock":
        run.font.color.rgb = RGBColor(0xAA, 0x00, 0x00)

    doc.add_paragraph(
        f"Карьер: {_fmt(data['quarry']['name'])} · Участок: {_fmt(data['quarry']['section'])} · "
        f"Взрыв: {_fmt(data['blast_event']['blast_datetime'])} · "
        f"Устройство: {_fmt(data['capture']['device'])}"
    )

    doc.add_heading("Результаты анализа", level=2)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, label in RESULT_FIELDS_RU:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = _fmt(data["analysis_result"].get(key))
    if data["analysis_result"]["confidence_notes"]:
        doc.add_paragraph(f"Примечания: {data['analysis_result']['confidence_notes']}")

    dist = data["analysis_result"]["size_distribution"]
    if dist:
        doc.add_heading("Гранулометрический состав", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        table.rows[0].cells[0].text = "Размер, мм"
        table.rows[0].cells[1].text = "Прохождение, %"
        for row in dist:
            cells = table.add_row().cells
            cells[0].text = _fmt(row.get("size_mm"))
            cells[1].text = _fmt(row.get("cumulative_passing_pct"))

    doc.add_heading("Паспорт БВР", level=2)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, label in PASSPORT_FIELDS_RU:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = _fmt(data["passport"].get(key))

    doc.add_heading("Рекомендации", level=2)
    if not data["recommendations"]:
        doc.add_paragraph("Рекомендаций нет.")
    for rec in data["recommendations"]:
        p = doc.add_paragraph()
        p.add_run(f"[{rec['status_ru']}] ").bold = True
        p.add_run(rec["recommendation_text"])
        if rec["parameter_suggestions"]:
            doc.add_paragraph(
                "Параметры (справочно, не применяются автоматически): "
                + json.dumps(rec["parameter_suggestions"], ensure_ascii=False)
            )
        if rec["reviewer_notes"]:
            doc.add_paragraph(f"Заметки проверяющего: {rec['reviewer_notes']}")

    image_titles = [
        ("left_frame", "Исходный снимок (левый кадр)"),
        ("mask_overlay", "Сегментация (наложение масок)"),
        ("depth_color", "Карта глубины"),
    ]
    if any(name in images for name, _ in image_titles):
        doc.add_heading("Снимки", level=2)
        for name, title in image_titles:
            if name in images:
                doc.add_paragraph(title)
                doc.add_picture(io.BytesIO(images[name]), width=Inches(6.0))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


_PDF_TEMPLATE = """
<!doctype html><html><head><meta charset="utf-8"><style>
  @page { size: A4; margin: 18mm 15mm; }
  body { font-family: "DejaVu Sans", sans-serif; font-size: 10pt; color: #1a1a1a; }
  h1 { font-size: 16pt; margin: 0 0 4pt 0; }
  h2 { font-size: 12pt; margin: 14pt 0 4pt 0; border-bottom: 1px solid #999; }
  .badge { font-weight: bold; padding: 4pt 8pt; border-radius: 4pt; display: inline-block; }
  .badge.mock { background: #fdecea; color: #a00; border: 1px solid #a00; }
  .badge.real { background: #e8f5e9; color: #1b5e20; border: 1px solid #1b5e20; }
  table { border-collapse: collapse; width: 100%; margin: 4pt 0; }
  td, th { border: 1px solid #bbb; padding: 3pt 6pt; text-align: left; vertical-align: top; }
  th { background: #f0f0f0; }
  .muted { color: #555; }
  img.visual { max-width: 100%; margin: 4pt 0; }
  .rec { margin: 6pt 0; padding: 6pt; border: 1px solid #ccc; border-radius: 4pt; }
  .rec .status { font-weight: bold; }
</style></head><body>
  <h1>{{ d.report.title }}</h1>
  <p><span class="badge {{ d.method.analysis_method }}">{{ d.method.label }}</span></p>
  <p class="muted">Сформирован: {{ d.exported_at }} · Карьер: {{ fmt(d.quarry.name) }} ·
     Участок: {{ fmt(d.quarry.section) }} · Взрыв: {{ fmt(d.blast_event.blast_datetime) }} ·
     Устройство: {{ fmt(d.capture.device) }}</p>

  <h2>Результаты анализа</h2>
  <table>
    {% for key, label in result_fields %}
    <tr><td>{{ label }}</td><td>{{ fmt(d.analysis_result[key]) }}</td></tr>
    {% endfor %}
  </table>
  {% if d.analysis_result.confidence_notes %}
  <p class="muted">{{ d.analysis_result.confidence_notes }}</p>
  {% endif %}

  {% if d.analysis_result.size_distribution %}
  <h2>Гранулометрический состав</h2>
  <table>
    <tr><th>Размер, мм</th><th>Прохождение, %</th></tr>
    {% for row in d.analysis_result.size_distribution %}
    <tr><td>{{ fmt(row.size_mm) }}</td><td>{{ fmt(row.cumulative_passing_pct) }}</td></tr>
    {% endfor %}
  </table>
  {% endif %}

  <h2>Паспорт БВР</h2>
  <table>
    {% for key, label in passport_fields %}
    <tr><td>{{ label }}</td><td>{{ fmt(d.passport[key]) }}</td></tr>
    {% endfor %}
  </table>

  <h2>Рекомендации</h2>
  {% if not d.recommendations %}<p>Рекомендаций нет.</p>{% endif %}
  {% for rec in d.recommendations %}
  <div class="rec">
    <div class="status">[{{ rec.status_ru }}]</div>
    <div style="white-space: pre-wrap">{{ rec.recommendation_text }}</div>
    {% if rec.parameter_suggestions %}
    <p class="muted">Параметры (справочно, не применяются автоматически):
       {{ rec.parameter_suggestions }}</p>
    {% endif %}
    {% if rec.reviewer_notes %}<p class="muted">Заметки: {{ rec.reviewer_notes }}</p>{% endif %}
  </div>
  {% endfor %}

  {% if image_tags %}
  <h2>Снимки</h2>
  {% for title, src in image_tags %}
  <p class="muted">{{ title }}</p>
  <img class="visual" src="{{ src }}"/>
  {% endfor %}
  {% endif %}
</body></html>
"""


def render_pdf(data: dict, images: dict[str, bytes] | None = None) -> bytes:
    import base64

    from jinja2 import Environment
    from weasyprint import HTML

    images = images or {}
    image_titles = [
        ("left_frame", "Исходный снимок (левый кадр)"),
        ("mask_overlay", "Сегментация (наложение масок)"),
        ("depth_color", "Карта глубины"),
    ]
    image_tags = [
        (title, "data:image/png;base64," + base64.b64encode(images[name]).decode())
        for name, title in image_titles
        if name in images
    ]

    env = Environment(autoescape=True)
    html = env.from_string(_PDF_TEMPLATE).render(
        d=_public(data),
        fmt=_fmt,
        result_fields=RESULT_FIELDS_RU,
        passport_fields=PASSPORT_FIELDS_RU,
        image_tags=image_tags,
    )
    return HTML(string=html).write_pdf()


def render(fmt: str, data: dict, images: dict[str, bytes] | None = None) -> bytes:
    if fmt == "json":
        return render_json(data)
    if fmt == "csv":
        return render_csv(data)
    if fmt == "xlsx":
        return render_xlsx(data)
    if fmt == "docx":
        return render_docx(data, images)
    if fmt == "pdf":
        return render_pdf(data, images)
    raise ValueError(f"Unsupported export format: {fmt}")
